#!/usr/bin/env python3
"""
skill-test engine.

Run a skill against a known input (its sample_spec.json) and check the output
structure before the skill is committed. Currently knows how to exercise the
document-producing skills (generate-avalere-pptx, generate-avalere-docx) and to
lint any SKILL.md frontmatter.

Usage:
    python3 skill_test.py <skill-dir-name|all>
"""
import sys, os, json, subprocess, tempfile, glob

SKILLS_DIR = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
REQUIRED_FRONTMATTER = ["name", "description", "triggers", "required_tools", "outputs", "sensitivity"]


def check_frontmatter(skill_dir):
    path = os.path.join(skill_dir, "SKILL.md")
    if not os.path.exists(path):
        return [f"missing SKILL.md"]
    txt = open(path, encoding="utf-8").read()
    if not txt.startswith("---"):
        return ["SKILL.md has no frontmatter block"]
    fm = txt.split("---", 2)[1]
    problems = [f"frontmatter missing '{k}'" for k in REQUIRED_FRONTMATTER if f"{k}:" not in fm]
    if "client:" not in fm:
        problems.append("frontmatter missing 'client' (client-portability rule: every skill is tenant-scoped)")
    return problems


def run_builder(skill_dir, builder, validator):
    spec = os.path.join(skill_dir, "sample_spec.json")
    if not os.path.exists(spec):
        return [f"no sample_spec.json to test against"]
    out = tempfile.mktemp(suffix=os.path.splitext(json.load(open(spec)).get("output", "out.bin"))[1])
    r = subprocess.run([sys.executable, os.path.join(skill_dir, builder), spec, "--out", out],
                       capture_output=True, text=True)
    if r.returncode != 0:
        return [f"builder failed: {r.stderr.strip()[:200]}"]
    if not os.path.exists(out) or os.path.getsize(out) == 0:
        return ["builder produced no output"]
    return validator(out)


def validate_pptx(path):
    from pptx import Presentation
    p = Presentation(path)
    n = len(list(p.slides))
    return [] if n > 0 else ["deck has no slides"]


def validate_docx(path):
    from docx import Document
    d = Document(path)
    n = len([p for p in d.paragraphs if p.text.strip()])
    return [] if n > 0 else ["document has no content paragraphs"]


BUILDERS = {
    "generate-avalere-pptx": ("build_deck.py", validate_pptx),
    "generate-avalere-docx": ("build_doc.py", validate_docx),
}


def test_skill(name):
    skill_dir = os.path.join(SKILLS_DIR, name)
    if not os.path.isdir(skill_dir):
        return [f"no such skill: {name}"]
    problems = check_frontmatter(skill_dir)
    if name in BUILDERS:
        builder, validator = BUILDERS[name]
        problems += run_builder(skill_dir, builder, validator)
    return problems


def main():
    if len(sys.argv) < 2:
        sys.exit("usage: skill_test.py <skill-dir-name|all>")
    target = sys.argv[1]
    names = ([os.path.basename(os.path.dirname(p)) for p in glob.glob(os.path.join(SKILLS_DIR, "*", "SKILL.md"))]
             if target == "all" else [target])
    failed = 0
    for name in sorted(names):
        problems = test_skill(name)
        if problems:
            failed += 1
            print(f"FAIL  {name}")
            for p in problems:
                print(f"      - {p}")
        else:
            print(f"PASS  {name}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
